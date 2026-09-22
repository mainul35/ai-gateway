"""Reading a document well enough for a model to answer questions about it.

A PDF, a Word file or a spreadsheet is not something a language model can be handed: it wants text.
So the text is taken out here and given to the model as part of the conversation, with the file's
name attached, and the model answers from that.

What comes out is plain and a little lossy. A spreadsheet becomes rows of tab-separated cells, a
Word file becomes its paragraphs and its tables, a PDF becomes whatever text was stored in it. That
is enough to answer questions about what a document says, and it is not enough to reproduce how it
looked, which is a distinction worth being honest about rather than hiding.

A PDF that holds no text at all is almost always a scan - pictures of pages. Nothing here can read
those, and saying so plainly is better than returning an empty string and letting a model invent an
answer from nothing.
"""
import io
import logging
import re

log = logging.getLogger("tools.documents")

# Enough of a document to answer questions about it, without burying the conversation
MAX_CHARACTERS = 60_000
MAX_PDF_PAGES = 80
MAX_SHEET_ROWS = 400

SUFFIXES = {
    ".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx", ".xlsm": "xlsx",
    ".txt": "text", ".md": "text", ".markdown": "text", ".csv": "text", ".tsv": "text",
    ".log": "text", ".json": "text", ".yaml": "text", ".yml": "text", ".xml": "text",
    ".html": "text", ".htm": "text", ".rst": "text", ".ini": "text", ".properties": "text",
}
# The old binary Word and Excel formats are a different thing entirely, and are not supported
OLD_OFFICE = (".doc", ".xls", ".ppt")

MIME = {"pdf": "application/pdf",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "text": "text/plain"}


class DocumentError(Exception):
    pass


def kind_of(name, data=b""):
    """Which reader this file needs, or None when it is not a document at all."""
    lowered = (name or "").lower()
    if lowered.endswith(OLD_OFFICE):
        raise DocumentError(f"{name} is in an older Office format. Save it as .docx, .xlsx or PDF "
                            "and attach that.")
    for suffix, kind in SUFFIXES.items():
        if lowered.endswith(suffix):
            return kind
    # Trust the bytes when the name says nothing: a PDF and a zip-based Office file both announce
    # themselves, and anything that decodes as text is text
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:2] == b"PK":
        return "docx" if b"word/" in data[:4000] else ("xlsx" if b"xl/" in data[:4000] else None)
    if data and _as_text(data) is not None:
        return "text"
    return None


def is_supported(name, data=b""):
    try:
        return kind_of(name, data) is not None
    except DocumentError:
        return False


def _as_text(data):
    """The bytes as text, or None when they are not text at all."""
    for encoding in ("utf-8", "utf-16", "cp1252"):
        try:
            decoded = data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
        # A binary file can decode as cp1252 and still be nonsense; control characters give it away
        controls = sum(1 for c in decoded[:2000] if ord(c) < 9 or 13 < ord(c) < 32)
        if controls > len(decoded[:2000]) * 0.02:
            return None
        return decoded
    return None


def _from_pdf(data):
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as e:
        raise DocumentError(f"This PDF could not be opened ({e}).")
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception:
            raise DocumentError("This PDF is password protected.")
    pages = []
    for number, page in enumerate(reader.pages[:MAX_PDF_PAGES], start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            pages.append(f"--- page {number} ---\n{text.strip()}")
    if not pages:
        raise DocumentError("This PDF has no text in it - it is almost certainly a scan. Nothing "
                            "here can read pictures of pages.")
    more = len(reader.pages) - MAX_PDF_PAGES
    if more > 0:
        pages.append(f"[and {more} more page(s), not read]")
    return "\n\n".join(pages)


def _from_docx(data):
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise DocumentError(f"This Word file could not be opened ({e}).")
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if any(row.strip() for row in rows):
            parts.append("\n".join(rows))
    if not parts:
        raise DocumentError("This Word file has no text in it.")
    return "\n\n".join(parts)


def _from_xlsx(data):
    import openpyxl

    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise DocumentError(f"This spreadsheet could not be opened ({e}).")
    sheets = []
    for sheet in book.worksheets:
        rows = []
        for number, row in enumerate(sheet.iter_rows(values_only=True)):
            if number >= MAX_SHEET_ROWS:
                rows.append(f"[more rows in {sheet.title}, not read]")
                break
            cells = ["" if c is None else str(c) for c in row]
            if any(cell.strip() for cell in cells):
                rows.append("\t".join(cells).rstrip())
        if rows:
            sheets.append(f"--- sheet: {sheet.title} ---\n" + "\n".join(rows))
    book.close()
    if not sheets:
        raise DocumentError("This spreadsheet is empty.")
    return "\n\n".join(sheets)


def extract(name, data):
    """The document as text. Raises DocumentError with something a person can act on."""
    kind = kind_of(name, data)
    if kind is None:
        raise DocumentError(f"{name} is not a document this server can read.")
    if kind == "pdf":
        text = _from_pdf(data)
    elif kind == "docx":
        text = _from_docx(data)
    elif kind == "xlsx":
        text = _from_xlsx(data)
    else:
        text = _as_text(data)
        if text is None:
            raise DocumentError(f"{name} is not readable as text.")
    text = re.sub(r"\n{4,}", "\n\n\n", text.replace("\r\n", "\n")).strip()
    if len(text) > MAX_CHARACTERS:
        text = text[:MAX_CHARACTERS] + "\n\n[the rest of this document was too long to include]"
    return text


def context_block(documents):
    """What the model is told about the files attached to this turn."""
    if not documents:
        return ""
    parts = ["The user has attached the following document(s). Answer from what they contain, and "
             "say so when something they ask about is not in them."]
    for name, text in documents:
        parts.append(f"\n===== {name} =====\n{text}")
    return "\n".join(parts)
